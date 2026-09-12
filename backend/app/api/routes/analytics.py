from collections import defaultdict
from datetime import datetime, timedelta, timezone

from fastapi import APIRouter, Depends, HTTPException, Query

from app.core.dependencies import get_repository
from app.repositories.memory import MemoryRepository

router = APIRouter(prefix="/analytics", tags=["analytics"])

# 统一时区口径：业务日期一律用东八区（库内时间戳为 UTC 存储）
TZ = timezone(timedelta(hours=8))


def _pareto(values: dict[str, int]) -> list[dict[str, int | str]]:
    total = sum(values.values()) or 1
    accumulated = 0
    result = []
    for name, count in sorted(values.items(), key=lambda item: item[1], reverse=True):
        accumulated += count
        result.append({"name": name, "count": count, "cumulative_rate": round(accumulated / total * 100)})
    return result


def _build_dimension_script(repo, leads: list) -> list[dict]:
    """P2-19b：按话术维度归因，基于 script_usages 统计触达/加微/成交/转化率"""
    usages = repo.list_script_usages()
    scripts = {s.id: s.name for s in repo.list_scripts()}
    lead_status = {l.id: l.status for l in leads}
    bucket: dict[str, dict[str, int]] = defaultdict(lambda: {"leads": 0, "wechat": 0, "won": 0})
    seen: dict[str, set] = defaultdict(set)
    for u in usages:
        key = u.script_id
        if u.lead_id in seen[key]:
            continue
        seen[key].add(u.lead_id)
        b = bucket[key]
        b["leads"] += 1
        st = lead_status.get(u.lead_id, "")
        if st == "wechat_added" or u.result == "wechat_added":
            b["wechat"] += 1
        if st == "deal_won" or u.result == "deal_won":
            b["won"] += 1
    rows = []
    for sid, b in bucket.items():
        rows.append({
            "name": scripts.get(sid, sid),
            "script_id": sid,
            "leads": b["leads"],
            "wechat": b["wechat"],
            "won": b["won"],
            "conversion_rate": round(b["won"] / b["leads"] * 100, 1) if b["leads"] else 0,
        })
    rows.sort(key=lambda x: x["leads"], reverse=True)
    return rows


def _build_dimension_account(repo, leads: list) -> list[dict]:
    """P2-19b：按账号维度归因，基于 lead.account 统计线索/加微/成交/转化率"""
    bucket: dict[str, dict[str, int]] = defaultdict(lambda: {"leads": 0, "wechat": 0, "won": 0})
    for lead in leads:
        key = lead.account or "未分配账号"
        b = bucket[key]
        b["leads"] += 1
        if lead.status == "wechat_added":
            b["wechat"] += 1
        if lead.status == "deal_won":
            b["won"] += 1
    rows = []
    for name, b in bucket.items():
        rows.append({
            "name": name,
            "leads": b["leads"],
            "wechat": b["wechat"],
            "won": b["won"],
            "conversion_rate": round(b["won"] / b["leads"] * 100, 1) if b["leads"] else 0,
        })
    rows.sort(key=lambda x: x["leads"], reverse=True)
    return rows


@router.get("/report")
def report(
    days: int = Query(30, ge=7, le=90),
    dimension: str = Query("source", description="归因维度：source / script / account"),
    repo: MemoryRepository = Depends(get_repository),
) -> dict:
    leads = repo.list_leads()
    events = repo.list_behavior_events()
    event_counts = defaultdict(int)
    for event in events:
        event_counts[event.event_type] += 1
    source: dict[str, dict[str, int | float]] = defaultdict(lambda: {"leads": 0, "high_intent": 0, "won": 0, "amount": 0})
    region: dict[str, dict[str, int]] = defaultdict(lambda: {"leads": 0, "high_intent": 0, "won": 0})
    daily: dict[str, int] = defaultdict(int)
    pain_points: dict[str, int] = defaultdict(int)
    for lead in leads:
        day = _local_date(lead.created_at)
        if day:
            daily[day] += 1
        pain_points[lead.customer_need or lead.note or "未记录需求"] += 1
        item = source[lead.source_keyword or "未标记"]
        item["leads"] += 1
        item["high_intent"] += lead.intent_level == "A"
        item["won"] += lead.status == "won"
        item["amount"] += lead.conversion_amount if lead.status == "won" else 0
        area = region[lead.region or "未填写"]
        area["leads"] += 1
        area["high_intent"] += lead.intent_level == "A"
        area["won"] += lead.status == "won"
    stages = {
        "leads": len(leads),
        "qualified": sum(lead.status in {"qualified", "contacted", "replied", "follow_up", "won"} for lead in leads),
        "replied": event_counts.get("reply", 0) or sum(lead.status in {"replied", "follow_up", "won"} for lead in leads),
        "follow_up": event_counts.get("follow_up", 0) or sum(lead.status in {"follow_up", "won"} for lead in leads),
        "won": event_counts.get("won", 0) or sum(lead.status == "won" for lead in leads),
    }
    events_by_lead: dict[str, list] = defaultdict(list)
    for event in events:
        events_by_lead[event.lead_id].append(event)
    rfm = [{
        "lead_id": lead.id,
        "nickname": lead.nickname,
        "recency": 3 if any(event.event_type in {"reply", "follow_up", "won"} for event in events_by_lead[lead.id]) else 1,
        "frequency": 3 if len(events_by_lead[lead.id]) >= 3 else 2 if len(events_by_lead[lead.id]) >= 2 else 1,
        "monetary": 3 if sum(event.value for event in events_by_lead[lead.id]) > 0 or lead.conversion_amount else 2 if lead.intent_level == "A" else 1,
        "segment": "高价值客户" if lead.status == "won" else "高潜待推进" if lead.intent_level == "A" else "持续培育" if lead.intent_level == "B" else "待激活",
    } for lead in sorted(leads, key=lambda item: (item.intent_level, item.conversion_amount), reverse=True)]
    def emotion_for(lead) -> str:
        text = f"{lead.note} {lead.customer_need}".lower()
        if any(token in text for token in ["想", "预算", "报价", "准备", "尽快", "确认"]):
            return "积极明确"
        if any(token in text for token in ["比较", "了解", "看看", "考虑"]):
            return "观望比较"
        return "低意向" if lead.intent_level not in {"A", "B"} else "积极明确"

    emotions = {key: sum(emotion_for(item) == key for item in leads) for key in ("积极明确", "观望比较", "低意向")}
    cohorts = []
    for day, size in sorted(daily.items()):
        members = [item for item in leads if _local_date(item.created_at) == day]
        member_events = {item.id: events_by_lead.get(item.id, []) for item in members}
        def retained_within(lead_id: str, limit: int) -> bool:
            events_for_lead = member_events[lead_id]
            return any(0 < (event.created_at - next(item.created_at for item in members if item.id == lead_id)).days <= limit for event in events_for_lead)
        cohorts.append({
            "date": day,
            "users": size,
            "d1": round(sum(retained_within(item.id, 1) for item in members) / size * 100) if events else round(sum(item.status != "new" for item in members) / size * 100) if size else 0,
            "d3": round(sum(retained_within(item.id, 3) for item in members) / size * 100) if events else round(sum(item.status in {"replied", "follow_up", "won"} for item in members) / size * 100) if size else 0,
            "d7": round(sum(retained_within(item.id, 7) for item in members) / size * 100) if events else round(sum(item.status == "won" for item in members) / size * 100) if size else 0,
            "basis": "基于用户行为事件" if events else "当前阶段代理留存",
        })
    daily_rows = [{"date": key, "leads": value} for key, value in sorted(daily.items())][-days:]
    cohort_rows = cohorts[-days:]
    top_source = max(source.items(), key=lambda item: item[1]["won"], default=("未标记", {"won": 0}))[0]
    top_pain = max(pain_points.items(), key=lambda item: item[1], default=("未记录需求", 0))[0]

    # P2-19b：dimension 参数切换归因分组
    dimension = (dimension or "source").lower()
    attribution_rows: list[dict] = []
    if dimension == "script":
        attribution_rows = _build_dimension_script(repo, leads)
    elif dimension == "account":
        attribution_rows = _build_dimension_account(repo, leads)
    # source 维度保持原有 sources 字段（向后兼容）

    return {
        "generated_at": datetime.now(timezone.utc),
        "summary": {"leads": len(leads), "high_intent": sum(x.intent_level == "A" for x in leads), "won_amount": sum(x.conversion_amount for x in leads if x.status == "won")},
        "period_days": days,
        "dimension": dimension,
        "daily": daily_rows,
        "funnel": stages,
        "sources": [{"name": key, **value, "high_intent_rate": round(value["high_intent"] / value["leads"] * 100) if value["leads"] else 0, "conversion_rate": round(value["won"] / value["leads"] * 100) if value["leads"] else 0} for key, value in sorted(source.items(), key=lambda item: item[1]["leads"], reverse=True)],
        "regions": [{"name": key, **value, "high_intent_rate": round(value["high_intent"] / value["leads"] * 100) if value["leads"] else 0, "conversion_rate": round(value["won"] / value["leads"] * 100) if value["leads"] else 0} for key, value in sorted(region.items(), key=lambda item: item[1]["leads"], reverse=True)],
        "attribution": attribution_rows,
        "rfm": rfm,
        "emotions": emotions,
        "emotion_basis": "基于用户原声关键词与意向等级兜底",
        "pain_points": _pareto(pain_points),
        "cohorts": cohort_rows,
        "insights": [
            f"当前有 {sum(item.intent_level == 'A' for item in leads)} 位 A 级客户需要优先跟进。",
            f"成交表现最好的来源是“{top_source}”，建议继续观察其转化质量。",
            f"用户反馈中出现最多的需求是“{top_pain}”，可用于下一轮内容和话术优化。",
        ],
        "path": [
            {"from": "内容线索", "to": "评论互动", "value": len(leads)},
            {"from": "评论互动", "to": "私信回复", "value": stages["replied"]},
            {"from": "私信回复", "to": "跟进", "value": stages["follow_up"]},
            {"from": "跟进", "to": "成交", "value": stages["won"]},
        ],
        "event_counts": dict(event_counts),
    }



# ============================================================
# P3-14: Ad spend entry endpoints
# ============================================================
from pydantic import BaseModel, Field
from uuid import uuid4


class SpendInput(BaseModel):
    date: str = Field(description="YYYY-MM-DD")
    platform: str = "douyin"
    amount: float = Field(gt=0)
    note: str = ""


@router.post("/spend")
def add_spend(payload: SpendInput, repo=Depends(get_repository)) -> dict:
    spend_id = uuid4().hex
    now = datetime.now(timezone.utc).isoformat()
    if hasattr(repo, "_execute"):
        repo._execute(
            "INSERT INTO ad_spend (id, date, platform, amount, note, created_at) VALUES (?, ?, ?, ?, ?, ?)",
            (spend_id, payload.date, payload.platform, payload.amount, payload.note, now),
        )
    return {"ok": True, "id": spend_id, "amount": payload.amount, "date": payload.date}


@router.get("/spend")
def list_spend(limit: int = Query(30, ge=1, le=365), repo=Depends(get_repository)) -> list:
    if hasattr(repo, "_query"):
        rows = repo._query(
            "SELECT id, date, platform, amount, note, created_at FROM ad_spend ORDER BY date DESC LIMIT ?",
            (limit,),
        )
        return [dict(r) for r in rows]
    return []


# ============================================================
# P3-15: 数据分析统一API（漏斗 / 维度分析 / 综合导出）
# ============================================================
from datetime import datetime, timedelta


def _local_date(value) -> str | None:
    """datetime / ISO 字符串 → 东八区日期 YYYY-MM-DD（统一日期归属口径）

    - 带 +00:00 / Z 偏移：按其自身时区解析后转东八区
    - 裸字符串（无时区）：按 UTC 解释（与库内 100% UTC 存储一致），不再叠加 +8
    """
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        dt = value
    else:
        s = str(value).strip()
        if not s:
            return None
        try:
            dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
        except ValueError:
            for fmt in ("%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
                try:
                    dt = datetime.strptime(s, fmt)
                    break
                except ValueError:
                    continue
            else:
                return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(TZ).date().isoformat()


def _dt_in_range(dt, start_date, end_date):
    """判断 datetime 的东八区日期是否在 [start_date, end_date] 范围内"""
    if dt is None:
        return False
    d = _local_date(dt)
    if d is None:
        return False
    if start_date and d < start_date:
        return False
    if end_date and d > end_date:
        return False
    return True


def _default_range():
    """默认近7天（按东八区当天，与过滤口径一致）"""
    today = datetime.now(TZ).date()
    start = today - timedelta(days=6)
    return start.isoformat(), today.isoformat()


# ============================================================
# BUG-1：日期参数格式校验（非法格式 / start>end → 422，不静默返回空）
# ============================================================
def _parse_date(value: str | None, field: str) -> str | None:
    """校验 YYYY-MM-DD；空串/None 视为未传（只回退该侧默认值，不影响另一侧）"""
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.strptime(text, "%Y-%m-%d").date().isoformat()
    except ValueError:
        raise HTTPException(
            status_code=422,
            detail=f"{field} 日期格式非法，必须为 YYYY-MM-DD（例如 2026-09-12），当前收到：{value!r}",
        )


def _resolve_range(start_date: str | None, end_date: str | None) -> tuple[str, str]:
    start = _parse_date(start_date, "start_date")
    end = _parse_date(end_date, "end_date")
    default_start, default_end = _default_range()
    start = start or default_start
    end = end or default_end
    if start > end:
        raise HTTPException(
            status_code=422,
            detail=f"start_date({start}) 不能晚于 end_date({end})",
        )
    return start, end


# ============================================================
# BUG-2 / BUG-3 / BUG-5：触达、回复去重口径 + 账号归属统一
# ============================================================
def _as_dt(value):
    """datetime / ISO 字符串 → 可比较的 aware datetime"""
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)


def _lead_message_stats(messages, conversations, start_date, end_date) -> tuple[set, set]:
    """触达/回复新口径（按 lead 去重）

    - 触达：被我方发过 out 消息的去重 lead 数（不是消息条数）
    - 回复：在我方首次 out **之后**才产生 in 消息的去重 lead 数
            （用户先开口、我方零出站的会话不计入回复）
    返回值保证 replied ⊆ contacted，因此 回复 ≤ 触达、回复率 ≤ 100。
    """
    conv_to_lead = {c.id: c.lead_id for c in conversations if getattr(c, "lead_id", None)}
    first_out: dict[str, datetime] = {}
    inbounds: dict[str, list] = defaultdict(list)
    for msg in messages:
        if not _dt_in_range(msg.created_at, start_date, end_date):
            continue
        lead_id = conv_to_lead.get(msg.conversation_id)
        ts = _as_dt(msg.created_at)
        if not lead_id or ts is None:
            continue
        if msg.direction == "out":
            if lead_id not in first_out or ts < first_out[lead_id]:
                first_out[lead_id] = ts
        elif msg.direction == "in":
            inbounds[lead_id].append(ts)
    contacted = set(first_out)
    replied = {
        lead_id for lead_id, ts_list in inbounds.items()
        if lead_id in first_out and any(ts > first_out[lead_id] for ts in ts_list)
    }
    return contacted, replied


_FUNNEL_SUGGESTIONS = {
    "曝光": "增加评论区采集量，扩大关键词覆盖和对标账号监控范围",
    "入库": "优化意向筛选规则，提高A/B级线索识别准确率，减少误筛",
    "触达": "提升私信发送量，检查账号健康度和发送频率限制",
    "回复": "优化首条私信话术，提高破冰率，缩短响应时间",
    "加微": "优化加微引导话术，在客户高意向节点及时推送企微名片",
    "成交": "加强跟进节奏，对A/B级线索设置自动提醒，推动报价和签约",
}


@router.get("/funnel")
def analytics_funnel(
    start_date: str | None = Query(None, description="开始日期 YYYY-MM-DD"),
    end_date: str | None = Query(None, description="结束日期 YYYY-MM-DD"),
    repo=Depends(get_repository),
) -> dict:
    """转化漏斗：曝光→入库→触达→回复→加微→成交，带瓶颈定位和优化建议"""
    start_date, end_date = _resolve_range(start_date, end_date)

    leads = repo.list_leads()
    messages = repo.list_messages()
    conversations = repo.list_conversations()
    crawl_tasks = repo.list_crawl_tasks()

    # BUG-4：曝光只取真实埋点（采集任务采集量），不再用入库线索数兜底
    exposure = sum(t.collected_count or 0 for t in crawl_tasks if _dt_in_range(t.created_at, start_date, end_date))
    exposure_available = exposure > 0
    exposure_note = "" if exposure_available else "未接入曝光埋点：曝光数与首级转化率暂不可计算，请勿以入库数代替曝光数"
    collected = len([l for l in leads if _dt_in_range(l.created_at, start_date, end_date)])
    # BUG-2：触达/回复改为按 lead 去重，且回复必须是触达之后的入站
    contacted_set, replied_set = _lead_message_stats(messages, conversations, start_date, end_date)
    contacted = len(contacted_set)
    replied = len(replied_set)
    wechat = len([l for l in leads if _dt_in_range(l.wechat_added_at, start_date, end_date)])
    deal = len([l for l in leads if l.status == "deal_won" and _dt_in_range(l.deal_at, start_date, end_date)])

    stages_raw = [
        ("曝光", exposure), ("入库", collected), ("触达", contacted),
        ("回复", replied), ("加微", wechat), ("成交", deal),
    ]

    stages = []
    bottleneck_idx = -1
    min_step_rate = float("inf")
    for i, (label, value) in enumerate(stages_raw):
        prev_value = stages_raw[i - 1][1] if i > 0 else None
        if i == 0:
            step_rate = None
        elif prev_value is None or prev_value <= 0:
            # 曝光无埋点 / 上一级为 0 时不做假转化率，显式置空
            step_rate = None
        else:
            step_rate = round(min(value / prev_value * 100, 100.0), 1)
        overall_rate = round(min(value / exposure * 100, 100.0), 1) if exposure > 0 else None
        stage = {"label": label, "value": value, "step_rate": step_rate, "overall_rate": overall_rate}
        if label == "曝光":
            stage["available"] = exposure_available
            stage["note"] = exposure_note
        stages.append(stage)
        if step_rate is not None and step_rate < min_step_rate:
            min_step_rate = step_rate
            bottleneck_idx = i

    bottleneck = stages[bottleneck_idx]["label"] if bottleneck_idx >= 0 else None
    suggestion = _FUNNEL_SUGGESTIONS.get(bottleneck, "") if bottleneck else ""

    return {
        "start_date": start_date, "end_date": end_date,
        "stages": stages, "bottleneck": bottleneck, "suggestion": suggestion,
        "exposure_available": exposure_available, "exposure_note": exposure_note,
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


@router.get("/by-account")
def analytics_by_account(
    start_date: str | None = Query(None),
    end_date: str | None = Query(None),
    repo=Depends(get_repository),
) -> dict:
    """按账号排名：触达/回复/加微数 + 健康度"""
    start_date, end_date = _resolve_range(start_date, end_date)

    leads = repo.list_leads()
    messages = repo.list_messages()
    conversations = repo.list_conversations()
    accounts = {a.name: a for a in repo.list_accounts()}

    # BUG-3：线索与消息统一账号归属。优先用 leads.account；
    # 为空时用 conversations.lead_id → conversations.account_name 回填，保证同一批客户落在同一账号桶
    lead_account = {l.id: (l.account or "") for l in leads}
    conv_account_by_lead: dict[str, str] = {}
    conv_to_lead: dict[str, str] = {}
    conv_to_account: dict[str, str] = {}
    for conv in conversations:
        conv_to_account[conv.id] = conv.account_name or "未分配"
        if conv.lead_id:
            conv_to_lead[conv.id] = conv.lead_id
            conv_account_by_lead[conv.lead_id] = conv.account_name or "未分配"

    def account_of_lead(lead_id: str) -> str:
        if lead_account.get(lead_id):
            return lead_account[lead_id]
        return conv_account_by_lead.get(lead_id) or "未分配"

    def account_of_conv(conv_id: str) -> str:
        lead_id = conv_to_lead.get(conv_id)
        if lead_id and lead_id in lead_account:
            return account_of_lead(lead_id)
        return conv_to_account.get(conv_id, "未分配")

    bucket = defaultdict(lambda: {"leads": 0, "contacted": 0, "replied": 0, "wechat": 0, "deal": 0})
    for lead in leads:
        if not _dt_in_range(lead.created_at, start_date, end_date):
            continue
        b = bucket[account_of_lead(lead.id)]
        b["leads"] += 1
        if _dt_in_range(lead.wechat_added_at, start_date, end_date):
            b["wechat"] += 1
        if lead.status == "deal_won" and _dt_in_range(lead.deal_at, start_date, end_date):
            b["deal"] += 1

    # BUG-2/BUG-5：触达、回复按 lead 去重后计入所属账号（replied ⊆ contacted）
    contacted_set, replied_set = _lead_message_stats(messages, conversations, start_date, end_date)
    # 一个 lead 只归属一个账号桶，避免同一客户被重复计数
    lead_to_bucket: dict[str, str] = {}
    for msg in messages:
        if msg.direction != "out" or not _dt_in_range(msg.created_at, start_date, end_date):
            continue
        lead_id = conv_to_lead.get(msg.conversation_id)
        if lead_id in contacted_set and lead_id not in lead_to_bucket:
            lead_to_bucket[lead_id] = account_of_conv(msg.conversation_id)
    for lead_id in contacted_set:
        bucket[lead_to_bucket.get(lead_id, "未分配")]["contacted"] += 1
    for lead_id in replied_set:
        bucket[lead_to_bucket.get(lead_id, "未分配")]["replied"] += 1

    rows = []
    for name, b in bucket.items():
        acc_info = accounts.get(name)
        health_score = acc_info.health_score if acc_info else 0
        # 分母为去重后的触达人数，replied ⊆ contacted ⇒ reply_rate 恒 ≤ 100
        reply_rate = round(min(b["replied"], b["contacted"]) / b["contacted"] * 100, 1) if b["contacted"] > 0 else 0.0
        wechat_rate = round(b["wechat"] / b["leads"] * 100, 1) if b["leads"] > 0 else 0.0
        rows.append({
            "name": name, "leads": b["leads"], "contacted": b["contacted"],
            "replied": b["replied"], "wechat": b["wechat"], "deal": b["deal"],
            "reply_rate": reply_rate, "wechat_rate": wechat_rate, "health_score": health_score,
        })
    rows.sort(key=lambda x: x["wechat"], reverse=True)
    return {"start_date": start_date, "end_date": end_date, "rows": rows}


@router.get("/by-script")
def analytics_by_script(
    start_date: str | None = Query(None),
    end_date: str | None = Query(None),
    repo=Depends(get_repository),
) -> dict:
    """按话术转化率：使用次数/转化率，TOP5和BOTTOM5"""
    start_date, end_date = _resolve_range(start_date, end_date)

    usages = repo.list_script_usages()
    scripts = {s.id: s.name for s in repo.list_scripts()}

    bucket = defaultdict(lambda: {"uses": 0, "leads": set(), "wechat": 0, "deal": 0})
    for u in usages:
        if not _dt_in_range(u.used_at, start_date, end_date):
            continue
        key = u.script_id or "未标记"
        b = bucket[key]
        b["uses"] += 1
        b["leads"].add(u.lead_id)
        if u.result == "wechat_added":
            b["wechat"] += 1
        elif u.result == "deal_won":
            b["deal"] += 1

    rows = []
    for sid, b in bucket.items():
        lead_count = len(b["leads"])
        conv_rate = round(b["wechat"] / lead_count * 100, 1) if lead_count > 0 else 0.0
        rows.append({
            "script_id": sid, "name": scripts.get(sid, sid),
            "uses": b["uses"], "leads": lead_count,
            "wechat": b["wechat"], "deal": b["deal"], "conversion_rate": conv_rate,
        })
    rows.sort(key=lambda x: x["uses"], reverse=True)

    eligible = [r for r in rows if r["uses"] >= 5]
    top5 = sorted(eligible, key=lambda x: x["conversion_rate"], reverse=True)[:5]
    bottom5 = sorted(eligible, key=lambda x: x["conversion_rate"])[:5]

    return {"start_date": start_date, "end_date": end_date, "rows": rows, "top5": top5, "bottom5": bottom5}


@router.get("/by-industry")
def analytics_by_industry(
    start_date: str | None = Query(None),
    end_date: str | None = Query(None),
    repo=Depends(get_repository),
) -> dict:
    """按行业效果：线索数/加微数/成交数（以来源关键词作为行业代理维度）"""
    start_date, end_date = _resolve_range(start_date, end_date)

    leads = repo.list_leads()
    bucket = defaultdict(lambda: {"leads": 0, "wechat": 0, "deal": 0, "amount": 0.0})
    for lead in leads:
        if not _dt_in_range(lead.created_at, start_date, end_date):
            continue
        key = lead.source_keyword or lead.customer_need or "未分类"
        b = bucket[key]
        b["leads"] += 1
        if _dt_in_range(lead.wechat_added_at, start_date, end_date):
            b["wechat"] += 1
        if lead.status == "deal_won" and _dt_in_range(lead.deal_at, start_date, end_date):
            b["deal"] += 1
            b["amount"] += float(lead.deal_amount or 0)

    rows = []
    for name, b in bucket.items():
        wechat_rate = round(b["wechat"] / b["leads"] * 100, 1) if b["leads"] > 0 else 0.0
        rows.append({
            "name": name, "leads": b["leads"], "wechat": b["wechat"],
            "deal": b["deal"], "deal_amount": round(b["amount"], 2), "wechat_rate": wechat_rate,
        })
    rows.sort(key=lambda x: x["leads"], reverse=True)
    return {"start_date": start_date, "end_date": end_date, "rows": rows}


@router.get("/by-source")
def analytics_by_source(
    start_date: str | None = Query(None),
    end_date: str | None = Query(None),
    repo=Depends(get_repository),
) -> dict:
    """按来源分析：各触点线索/加微/成交/转化率"""
    start_date, end_date = _resolve_range(start_date, end_date)

    source_labels = {
        "own_comment": "自有视频评论区", "competitor": "对标账号监控",
        "inbound_dm": "用户主动私信", "fan_dm": "粉丝私信",
        "lead_card": "留资卡", "referral": "老客转介绍",
    }
    leads = repo.list_leads()
    bucket = defaultdict(lambda: {"leads": 0, "wechat": 0, "deal": 0, "amount": 0.0})
    for lead in leads:
        if not _dt_in_range(lead.created_at, start_date, end_date):
            continue
        key = lead.source or "unknown"
        b = bucket[key]
        b["leads"] += 1
        if _dt_in_range(lead.wechat_added_at, start_date, end_date):
            b["wechat"] += 1
        if lead.status == "deal_won" and _dt_in_range(lead.deal_at, start_date, end_date):
            b["deal"] += 1
            b["amount"] += float(lead.deal_amount or 0)

    rows = []
    for src, b in bucket.items():
        wechat_rate = round(b["wechat"] / b["leads"] * 100, 1) if b["leads"] > 0 else 0.0
        deal_rate = round(b["deal"] / b["leads"] * 100, 1) if b["leads"] > 0 else 0.0
        rows.append({
            "source": src, "label": source_labels.get(src, src),
            "leads": b["leads"], "wechat": b["wechat"], "deal": b["deal"],
            "deal_amount": round(b["amount"], 2), "wechat_rate": wechat_rate, "deal_rate": deal_rate,
        })
    rows.sort(key=lambda x: x["leads"], reverse=True)
    return {"start_date": start_date, "end_date": end_date, "rows": rows}


@router.get("/export")
def analytics_export(
    format: str = Query("xlsx", pattern="^(csv|xlsx)$"),
    start_date: str | None = Query(None),
    end_date: str | None = Query(None),
    repo=Depends(get_repository),
) -> dict:
    """综合导出：当前时间范围所有数据（线索/客户/成交），xlsx多sheet"""
    from app.services.export_service import ExportService, EXPORT_DIR, _ensure_export_dir, _gen_filename, _serialize_value, _try_openpyxl_available
    import csv as _csv

    start_date, end_date = _resolve_range(start_date, end_date)

    leads = [l for l in repo.list_leads() if _dt_in_range(l.created_at, start_date, end_date)]
    customers = [c for c in repo.list_customers() if _dt_in_range(c.created_at, start_date, end_date)]
    deals = [l for l in repo.list_leads() if l.status == "deal_won" and _dt_in_range(l.deal_at, start_date, end_date)]

    lead_cols = [("昵称", "nickname"), ("平台", "platform"), ("来源", "source"),
                 ("来源关键词", "source_keyword"), ("意向等级", "intent_level"),
                 ("状态", "status"), ("地区", "region"), ("客户需求", "customer_need"),
                 ("加微时间", "wechat_added_at"), ("成交金额", "deal_amount"),
                 ("成交时间", "deal_at"), ("创建时间", "created_at")]
    customer_cols = [("客户名称", "name"), ("来源", "source"), ("阶段", "stage"),
                     ("预估价值", "est_value"), ("成交金额", "deal_amount"),
                     ("加微时间", "wechat_added_at"), ("创建时间", "created_at")]
    deal_cols = [("客户昵称", "nickname"), ("来源", "source"), ("成交金额", "deal_amount"),
                 ("成交时间", "deal_at"), ("意向等级", "intent_level"), ("地区", "region")]

    use_xlsx = format.lower() == "xlsx" and _try_openpyxl_available()
    actual_format = "xlsx" if use_xlsx else "csv"
    ext = "xlsx" if use_xlsx else "csv"
    file_name = _gen_filename(f"analytics_{start_date}_{end_date}", ext)
    file_path = EXPORT_DIR / file_name
    _ensure_export_dir()

    if use_xlsx:
        import openpyxl
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "线索"
        ws1.append([c[0] for c in lead_cols])
        for l in leads:
            ws1.append([_serialize_value(getattr(l, f, "")) for _, f in lead_cols])
        ws2 = wb.create_sheet("客户")
        ws2.append([c[0] for c in customer_cols])
        for c in customers:
            ws2.append([_serialize_value(getattr(c, f, "")) for _, f in customer_cols])
        ws3 = wb.create_sheet("成交")
        ws3.append([c[0] for c in deal_cols])
        for d in deals:
            ws3.append([_serialize_value(getattr(d, f, "")) for _, f in deal_cols])
        for ws in (ws1, ws2, ws3):
            for col_idx in range(1, ws.max_column + 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 18
        wb.save(str(file_path))
    else:
        with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = _csv.writer(f)
            writer.writerow([c[0] for c in lead_cols])
            for l in leads:
                writer.writerow([_serialize_value(getattr(l, f, "")) for _, f in lead_cols])

    return {
        "export_id": uuid4().hex,
        "file_url": f"/exports/{file_name}",
        "file_name": file_name,
        "row_count": len(leads) + len(customers) + len(deals),
        "format": actual_format,
        "sheets": {"leads": len(leads), "customers": len(customers), "deals": len(deals)},
        "start_date": start_date, "end_date": end_date,
    }
