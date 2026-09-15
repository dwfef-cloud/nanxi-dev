# -*- coding: utf-8 -*-
"""解析 测试用例全集.md，生成可勾选的 Excel 工作簿。"""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

SRC = r"D:\nanxi-dev\测试用例全集.md"
OUT = r"D:\nanxi-dev\测试用例全集.xlsx"

# ---------- 1. 解析 markdown ----------
with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

cases = []
module = ""
guide_body = []      # 一、执行说明 正文
regress_body = []    # 二十二、通过标准 正文
capture = None

for ln in lines:
    m_sec = re.match(r'^##\s+\d+[、.]\s*(.*)$', ln)
    if m_sec:
        sec_title = m_sec.group(1)
        mm = re.match(r'^([^（]+)', sec_title)
        module = (mm.group(1) if mm else sec_title).strip().rstrip(' /').strip()
        if '执行说明' in sec_title:
            capture = 'guide'
        elif '通过标准' in sec_title:
            capture = 'regress'
        else:
            capture = None
        continue

    m_case = re.match(r'^###\s+(TC-[\w]+-\d+)\s+(.*)$', ln)
    if m_case:
        capture = None
        cid = m_case.group(1)
        rest = m_case.group(2)
        mp = re.search(r'【(P\d)·([A-Z]+)】', rest)
        prio = mp.group(1) if mp else ''
        ctype = mp.group(2) if mp else ''
        title = re.split(r'【', rest)[0].strip()
        mt = re.search(r'★\S+', rest)
        tag = mt.group(0).strip() if mt else ''
        cases.append({
            'module': module, 'id': cid, 'title': title,
            'prio': prio, 'type': ctype, 'tag': tag,
            'pre': '', 'steps': '', 'expect': ''
        })
        continue

    if capture in ('guide', 'regress'):
        if ln.strip():
            (guide_body if capture == 'guide' else regress_body).append(ln.strip())
        continue

    if cases:
        cur = cases[-1]
        if re.match(r'^\s*-\s*前置[：:]', ln):
            cur['pre'] = re.sub(r'^\s*-\s*前置[：:]\s*', '', ln).strip()
            continue
        if re.match(r'^\s*-\s*步骤[：:]', ln):
            cur['steps'] = re.sub(r'^\s*-\s*步骤[：:]\s*', '', ln).strip()
            continue
        if re.match(r'^\s+\d+\.\s', ln):
            cur['steps'] = (cur['steps'] + '\n' + ln.strip()).strip()
            continue
        if re.match(r'^\s*-\s*预期[：:]', ln):
            cur['expect'] = re.sub(r'^\s*-\s*预期[：:]\s*', '', ln).strip()
            continue

print('parsed cases:', len(cases))

# ---------- 2. 样式 ----------
def xl_color(css): return "FF" + css.removeprefix("#").upper()
XL_HEADER = xl_color("#4472C4")
XL_INPUT  = xl_color("#D9E2F3")
XL_TITLE  = xl_color("#2F5597")
XL_P0 = xl_color("#FFC7CE"); FT_P0 = xl_color("#9C0006")
XL_P1 = xl_color("#FFEB9C"); FT_P1 = xl_color("#9C6500")
XL_P2 = xl_color("#D9E2F3"); FT_P2 = xl_color("#1F3864")
XL_PASS = xl_color("#C6EFCE"); FT_PASS = xl_color("#006100")
XL_FAIL = xl_color("#FFC7CE"); FT_FAIL = xl_color("#9C0006")
XL_BLOCK= xl_color("#FFEB9C"); FT_BLOCK= xl_color("#9C6500")

thin = Side(style="thin", color="FFBFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_fill = PatternFill("solid", fgColor=XL_HEADER)
header_font = Font(bold=True, color="FFFFFFFF", size=11)
title_font = Font(bold=True, color="FFFFFFFF", size=14)
title_fill = PatternFill("solid", fgColor=XL_TITLE)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="top", wrap_text=True)

wb = Workbook()

# ---------- 3. 执行统计 sheet ----------
ws_s = wb.active
ws_s.title = "执行统计"
last = len(cases) + 1  # 明细数据最后一行（标题行1，数据2..last）

ws_s.merge_cells("A1:C1")
ws_s["A1"] = "获客系统 · 测试用例执行统计"
ws_s["A1"].font = title_font; ws_s["A1"].fill = title_fill
ws_s["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_s.row_dimensions[1].height = 28

stat_rows = [
    ("总用例数", f"=COUNTA(测试用例!B2:B{last})"),
    ("P0 数量",  f'=COUNTIF(测试用例!E2:E{last},"P0")'),
    ("P1 数量",  f'=COUNTIF(测试用例!E2:E{last},"P1")'),
    ("P2 数量",  f'=COUNTIF(测试用例!E2:E{last},"P2")'),
    ("通过",     f'=COUNTIF(测试用例!K2:K{last},"通过")'),
    ("失败",     f'=COUNTIF(测试用例!K2:K{last},"失败")'),
    ("阻塞",     f'=COUNTIF(测试用例!K2:K{last},"阻塞")'),
    ("未执行",   f"=B3-(B7+B8+B9)"),
    ("通过率",   f'=IF((B7+B8+B9)=0,"",B7/(B7+B8+B9))'),
]
r = 3
for name, formula in stat_rows:
    ws_s[f"A{r}"] = name
    ws_s[f"A{r}"].font = Font(bold=True)
    ws_s[f"B{r}"] = formula
    ws_s[f"A{r}"].border = border; ws_s[f"B{r}"].border = border
    ws_s[f"A{r}"].alignment = left
    if name == "通过率":
        ws_s[f"B{r}"].number_format = "0.0%"
    r += 1
ws_s[f"B{r-1}"].font = Font(bold=True, color=FT_PASS if False else "FF1F3864")

# 模块分布
r += 1
ws_s[f"A{r}"] = "模块分布"; ws_s[f"A{r}"].font = Font(bold=True, size=12); r += 1
ws_s[f"A{r}"] = "模块"; ws_s[f"B{r}"] = "用例数"
ws_s[f"A{r}"].font = header_font; ws_s[f"A{r}"].fill = header_fill
ws_s[f"B{r}"].font = header_font; ws_s[f"B{r}"].fill = header_fill
ws_s[f"A{r}"].alignment = center; ws_s[f"B{r}"].alignment = center
ws_s[f"A{r}"].border = border; ws_s[f"B{r}"].border = border
mod_start = r + 1
seen = []
for c in cases:
    if c['module'] not in seen:
        seen.append(c['module'])
r += 1
for mod in seen:
    ws_s[f"A{r}"] = mod
    ws_s[f"B{r}"] = f'=COUNTIF(测试用例!C2:C{last},"{mod}")'
    ws_s[f"A{r}"].border = border; ws_s[f"B{r}"].border = border
    ws_s[f"A{r}"].alignment = left
    r += 1
ws_s[f"A{r}"] = "合计"; ws_s[f"A{r}"].font = Font(bold=True)
ws_s[f"B{r}"] = f"=SUM(B{mod_start}:B{r-1})"
ws_s[f"A{r}"].border = border; ws_s[f"B{r}"].border = border

ws_s.column_dimensions["A"].width = 16
ws_s.column_dimensions["B"].width = 14
ws_s.column_dimensions["C"].width = 14

# ---------- 4. 测试用例明细 sheet ----------
ws = wb.create_sheet("测试用例")
headers = ["序号", "用例编号", "模块", "标题", "优先级", "类型", "标记",
           "前置条件", "测试步骤", "预期结果", "执行结果", "实际现象", "执行人", "执行日期"]
for ci, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=ci, value=h)
    c.fill = header_fill; c.font = header_font; c.alignment = center; c.border = border

wrap_cols = {3, 4, 7, 8, 9, 10, 12}
for i, c in enumerate(cases, start=2):
    row = [
        i - 1, c['id'], c['module'], c['title'], c['prio'], c['type'], c['tag'],
        c['pre'], c['steps'], c['expect'], "", "", "", ""
    ]
    for ci, val in enumerate(row, 1):
        cell = ws.cell(row=i, column=ci, value=val)
        cell.border = border
        if ci in wrap_cols:
            cell.alignment = left
        else:
            cell.alignment = center
        if ci in (5, 6, 11):  # 优先级/类型/结果 居中
            cell.alignment = center

last_row = len(cases) + 1

# 列宽
widths = {1:6, 2:14, 3:14, 4:28, 5:8, 6:8, 7:12, 8:30, 9:42, 10:48, 11:10, 12:30, 13:10, 14:14}
for col, w in widths.items():
    ws.column_dimensions[get_column_letter(col)].width = w

# 冻结表头 + 首列，开启筛选
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:N{last_row}"

# 数据验证下拉
dv_prio = DataValidation(type="list", formula1='"P0,P1,P2"', allow_blank=True)
dv_type = DataValidation(type="list", formula1='"UI,API,E2E,ERR"', allow_blank=True)
dv_res  = DataValidation(type="list", formula1='"通过,失败,阻塞,未执行"', allow_blank=True)
ws.add_data_validation(dv_prio); ws.add_data_validation(dv_type); ws.add_data_validation(dv_res)
dv_prio.add(f"E2:E{last_row}"); dv_type.add(f"F2:F{last_row}"); dv_res.add(f"K2:K{last_row}")

# 条件格式：优先级
ws.conditional_formatting.add(f"E2:E{last_row}",
    CellIsRule(operator="equal", formula=['"P0"'], fill=PatternFill("solid", fgColor=XL_P0), font=Font(color=FT_P0, bold=True)))
ws.conditional_formatting.add(f"E2:E{last_row}",
    CellIsRule(operator="equal", formula=['"P1"'], fill=PatternFill("solid", fgColor=XL_P1), font=Font(color=FT_P1)))
ws.conditional_formatting.add(f"E2:E{last_row}",
    CellIsRule(operator="equal", formula=['"P2"'], fill=PatternFill("solid", fgColor=XL_P2), font=Font(color=FT_P2)))
# 条件格式：执行结果
ws.conditional_formatting.add(f"K2:K{last_row}",
    CellIsRule(operator="equal", formula=['"通过"'], fill=PatternFill("solid", fgColor=XL_PASS), font=Font(color=FT_PASS, bold=True)))
ws.conditional_formatting.add(f"K2:K{last_row}",
    CellIsRule(operator="equal", formula=['"失败"'], fill=PatternFill("solid", fgColor=XL_FAIL), font=Font(color=FT_FAIL, bold=True)))
ws.conditional_formatting.add(f"K2:K{last_row}",
    CellIsRule(operator="equal", formula=['"阻塞"'], fill=PatternFill("solid", fgColor=XL_BLOCK), font=Font(color=FT_BLOCK, bold=True)))

# ---------- 5. 说明 sheet ----------
ws_i = wb.create_sheet("说明")
ws_i.merge_cells("A1:A1")
ws_i["A1"] = "使用说明 / 启动 / 回归清单"
ws_i["A1"].font = title_font; ws_i["A1"].fill = title_fill
ws_i["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws_i.row_dimensions[1].height = 24
r = 3
def put_block(title, body_lines):
    global r
    ws_i.cell(row=r, column=1, value=title).font = Font(bold=True, size=12, color="FF2F5597")
    r += 1
    for bl in body_lines:
        cell = ws_i.cell(row=r, column=1, value=bl)
        cell.alignment = left
        r += 1
    r += 1

if guide_body:
    put_block("一、执行说明（原文）", guide_body)
if regress_body:
    put_block("二十二、通过标准与回归建议（原文）", regress_body)

ws_i.column_dimensions["A"].width = 120

# ---------- 6. 标题属性 & 保存 ----------
wb.properties.title = "获客系统测试用例全集"
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved:", OUT, "cases:", len(cases))
